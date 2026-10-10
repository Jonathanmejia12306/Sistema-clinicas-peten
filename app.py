import sqlite3
import os
import uuid
import time
import io
from flask import Flask, render_template, g, request, redirect, url_for, flash, jsonify, send_from_directory, send_file
from datetime import date, datetime, timedelta
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from functools import wraps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# --- 0. CONFIGURACIÓN DE ZONA HORARIA (GUATEMALA) ---
# =============================================================================
os.environ['TZ'] = 'America/Guatemala'
try:
    time.tzset()
except AttributeError:
    pass

# =============================================================================
# --- 1. CONFIGURACIÓN DE LA APLICACIÓN ---
# =============================================================================

DATABASE = 'clinica.db'
UPLOAD_FOLDER = 'uploads'
PRODUCT_UPLOAD_FOLDER = os.path.join(UPLOAD_FOLDER, 'productos')
EXPENSE_UPLOAD_FOLDER = os.path.join(UPLOAD_FOLDER, 'comprobantes_gastos')

app = Flask(__name__)
app.config['SECRET_KEY'] = 'clave_maestra_clinica_46663627'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['PRODUCT_UPLOAD_FOLDER'] = PRODUCT_UPLOAD_FOLDER
app.config['EXPENSE_UPLOAD_FOLDER'] = EXPENSE_UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = None

# --- CONFIGURACIÓN DE LOGIN ---
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

# =============================================================================
# --- 2. SISTEMA DE USUARIOS Y PERMISOS ---
# =============================================================================

class User(UserMixin):
    def __init__(self, id, username, nombre, rol):
        self.id = id
        self.username = username
        self.nombre = nombre
        self.rol = rol

@login_manager.user_loader
def load_user(user_id):
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    user_db = conn.execute("SELECT * FROM usuarios WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    if user_db:
        return User(user_db['id'], user_db['username'], user_db['nombre'], user_db['rol'])
    return None

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.rol != 'admin':
            flash('Acceso denegado: Se requieren permisos de administrador.', 'error')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        conn = sqlite3.connect(DATABASE)
        conn.row_factory = sqlite3.Row
        user_db = conn.execute("SELECT * FROM usuarios WHERE username = ?", (username,)).fetchone()
        conn.close()
        if user_db and check_password_hash(user_db['password'], password):
            user_obj = User(user_db['id'], user_db['username'], user_db['nombre'], user_db['rol'])
            login_user(user_obj)
            return redirect(url_for('dashboard'))
        else:
            flash('Usuario o contraseña incorrectos', 'error')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

# --- GESTIÓN DE USUARIOS ---
@app.route('/usuarios')
@login_required
@admin_required
def lista_usuarios():
    db = get_db()
    usuarios = db.execute("SELECT id, username, nombre, rol FROM usuarios WHERE username != 'admin'").fetchall()
    return render_template('usuarios.html', active_page='usuarios', usuarios=usuarios)

@app.route('/usuarios/nuevo', methods=['POST'])
@login_required
@admin_required
def nuevo_usuario():
    username = request.form.get('username')
    password = request.form.get('password')
    nombre = request.form.get('nombre')
    rol = request.form.get('rol')
    hashed_pw = generate_password_hash(password)
    db = get_db()
    try:
        db.execute("INSERT INTO usuarios (username, password, nombre, rol) VALUES (?, ?, ?, ?)",
                   (username, hashed_pw, nombre, rol))
        db.commit()
        flash("Usuario creado con éxito.", "success")
    except sqlite3.IntegrityError:
        flash("El nombre de usuario ya existe.", "error")
    return redirect(url_for('lista_usuarios'))

@app.route('/usuarios/eliminar/<int:id>', methods=['POST'])
@login_required
@admin_required
def eliminar_usuario(id):
    db = get_db()
    db.execute("DELETE FROM usuarios WHERE id = ?", (id,))
    db.commit()
    flash("Usuario eliminado correctamente.", "success")
    return redirect(url_for('lista_usuarios'))

# =============================================================================
# --- 3. FUNCIONES AUXILIARES Y GESTIÓN DE ARCHIVOS ---
# =============================================================================

def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(exception):
    db = g.pop('db', None)
    if db is not None:
        db.close()

def calcular_edad(f):
    if not f: return None
    try:
        fn = datetime.strptime(f, '%Y-%m-%d').date()
        hoy = date.today()
        return hoy.year - fn.year - ((hoy.month, hoy.day) < (fn.month, fn.day))
    except: return None

def allowed_file(filename):
    return filename != ''

def save_unique_file(file, folder):
    original_filename = secure_filename(file.filename)
    extension = original_filename.rsplit('.', 1)[1].lower() if '.' in original_filename else ''
    unique_filename = f"{uuid.uuid4().hex}_{original_filename}"
    file.save(os.path.join(folder, unique_filename))
    return unique_filename

def get_caja_abierta(db):
    try:
        return db.execute("SELECT * FROM cierres_caja WHERE estado = 'Abierta'").fetchone()
    except sqlite3.OperationalError:
        return None

def procesar_formulario_inventario(conn, item_id=None):
    nombre = request.form.get('nombre')
    tipo = request.form.get('tipo')
    descripcion = request.form.get('descripcion', '')
    codigo_barras = request.form.get('codigo_barras', '')
    precio_venta = request.form.get('precio_venta', 0, type=float)
    precio_compra = request.form.get('precio_compra')
    porcentaje_ganancia = request.form.get('porcentaje_ganancia')
    cantidad = request.form.get('cantidad')
    fecha_vencimiento = request.form.get('fecha_vencimiento')
    ruta_foto = request.form.get('ruta_foto_actual')

    if 'foto' in request.files:
        file = request.files['foto']
        if file.filename != '':
            ruta_foto = save_unique_file(file, app.config['PRODUCT_UPLOAD_FOLDER'])

    precio_compra = float(precio_compra) if precio_compra else None
    porcentaje_ganancia = float(porcentaje_ganancia) if porcentaje_ganancia else None
    fecha_vencimiento = fecha_vencimiento if fecha_vencimiento else None

    # CORRECCIÓN: Manejo robusto de cantidad para evitar que se guarde como texto vacío
    try:
        cant_val = int(cantidad) if (cantidad and str(cantidad).strip() != "") else 0
    except (ValueError, TypeError):
        cant_val = 0

    if tipo == 'Servicio':
        precio_compra = porcentaje_ganancia = fecha_vencimiento = ruta_foto = None
        cant_val = None

    if item_id:
        conn.execute("UPDATE inventario SET nombre=?, descripcion=?, tipo=?, codigo_barras=?, precio_venta=?, precio_compra=?, porcentaje_ganancia=?, cantidad=?, fecha_vencimiento=?, ruta_foto=? WHERE id=?",
                     (nombre, descripcion, tipo, codigo_barras, precio_venta, precio_compra, porcentaje_ganancia, cant_val, fecha_vencimiento, ruta_foto, item_id))
    else:
        conn.execute("INSERT INTO inventario (nombre, descripcion, tipo, codigo_barras, precio_venta, precio_compra, porcentaje_ganancia, cantidad, fecha_vencimiento, ruta_foto) VALUES (?,?,?,?,?,?,?,?,?,?)",
                     (nombre, descripcion, tipo, codigo_barras, precio_venta, precio_compra, porcentaje_ganancia, cant_val, fecha_vencimiento, ruta_foto))
    conn.commit()

# =============================================================================
# --- 4. RUTA DASHBOARD ---
# =============================================================================

@app.route('/')
@login_required
def dashboard():
    conn = get_db()
    stats = {}
    try:
        stats['total_pacientes'] = conn.execute('SELECT COUNT(id) FROM pacientes WHERE is_deleted = 0').fetchone()[0]
        hoy = date.today().strftime('%Y-%m-%d')
        ventas = conn.execute("SELECT SUM(monto_total) FROM ventas WHERE date(fecha_venta) = ? AND estado != 'Cancelada'", (hoy,)).fetchone()[0]
        stats['ventas_hoy'] = round(ventas, 2) if ventas else 0
        stats['bajo_stock'] = conn.execute("SELECT COUNT(id) FROM inventario WHERE tipo='Producto' AND cantidad < 5 AND is_deleted = 0").fetchone()[0]
        fecha_limite = (date.today() + timedelta(days=90)).strftime('%Y-%m-%d')
        stats['por_vencer'] = conn.execute("SELECT COUNT(id) FROM inventario WHERE tipo='Producto' AND fecha_vencimiento BETWEEN ? AND ? AND is_deleted = 0", (hoy, fecha_limite)).fetchone()[0]
        proximas_citas = conn.execute("SELECT c.id, c.fecha_cita, c.hora_cita, c.descripcion, p.nombre, p.apellido FROM citas c JOIN pacientes p ON c.id_paciente = p.id WHERE c.fecha_cita >= ? AND c.estado = 'Pendiente' ORDER BY c.fecha_cita ASC, c.hora_cita ASC LIMIT 10", (hoy,)).fetchall()

        inv_stats = conn.execute("""
            SELECT
                SUM(COALESCE(cantidad, 0) * COALESCE(precio_compra, 0)),
                SUM(COALESCE(cantidad, 0) * COALESCE(precio_venta, 0))
            FROM inventario
            WHERE tipo = 'Producto' AND is_deleted = 0
        """).fetchone()
        costo_total = inv_stats[0] if inv_stats[0] else 0
        venta_total = inv_stats[1] if inv_stats[1] else 0
        stats['inv_costo'] = costo_total
        stats['inv_venta'] = venta_total
        stats['inv_ganancia'] = venta_total - costo_total
    except Exception as e:
        stats = {'total_pacientes': 0, 'ventas_hoy': 0, 'bajo_stock': 0, 'por_vencer': 0, 'inv_costo': 0, 'inv_venta': 0, 'inv_ganancia': 0}
        proximas_citas = []
    return render_template('index.html', active_page='dashboard', stats=stats, proximas_citas=proximas_citas)

# =============================================================================
# --- 5. APIs (JSON ENDPOINTS) ---
# =============================================================================

@app.route('/api/pacientes/buscar')
@login_required
def buscar_pacientes():
    q = f"%{request.args.get('q', '', type=str)}%"
    db = get_db()
    pacientes_db = db.execute("SELECT * FROM pacientes WHERE (nombre LIKE ? OR apellido LIKE ? OR (nombre || ' ' || apellido) LIKE ? OR telefono LIKE ?) AND is_deleted = 0 ORDER BY nombre", (q,q,q,q)).fetchall()
    return jsonify([dict(p, edad=calcular_edad(p['fecha_nacimiento'])) for p in pacientes_db])

@app.route('/api/expediente/buscar/<int:id_paciente>')
@login_required
def buscar_expedientes(id_paciente):
    q = f"%{request.args.get('q', '', type=str)}%"
    db = get_db()
    expedientes_db = db.execute("SELECT * FROM expedientes WHERE id_paciente = ? AND (CAST(id_expediente AS TEXT) LIKE ? OR fecha_expediente LIKE ? OR motivo_consulta LIKE ?) ORDER BY fecha_expediente DESC", (id_paciente, q, q, q)).fetchall()
    return jsonify([dict(row) for row in expedientes_db])

@app.route('/api/inventario/buscar')
@login_required
def api_inventario_buscar():
    query = request.args.get('q', '', type=str)
    filtro_costo_cero = request.args.get('costo_cero') == 'true'
    filtro_bajo_stock = request.args.get('bajo_stock') == 'true'
    filtro_por_vencer = request.args.get('por_vencer') == 'true'
    conn = get_db()
    base_sql = "SELECT * FROM inventario"
    conditions = ["is_deleted = 0"]
    params = []
    if query:
        conditions.append("(nombre LIKE ? OR tipo LIKE ? OR codigo_barras LIKE ?)")
        params.extend([f"%{query}%", f"%{query}%", f"%{query}%"])
    if filtro_costo_cero:
        conditions.append("(tipo = 'Producto' AND (precio_compra IS NULL OR precio_compra = 0))")
    if filtro_bajo_stock:
        conditions.append("(tipo = 'Producto' AND cantidad < 5)")
    if filtro_por_vencer:
        hoy = date.today().strftime('%Y-%m-%d'); fecha_limite = (date.today() + timedelta(days=90)).strftime('%Y-%m-%d')
        conditions.append("(fecha_vencimiento BETWEEN ? AND ?)"); params.extend([hoy, fecha_limite])
    order_by = "ORDER BY tipo DESC, nombre"
    if filtro_por_vencer:
        order_by = "ORDER BY fecha_vencimiento ASC, nombre"
    sql = f"{base_sql} WHERE {' AND '.join(conditions)} {order_by}"
    items_db = conn.execute(sql, tuple(params)).fetchall()
    items = []
    for item in items_db:
        item_dict = dict(item)
        precio_venta = item_dict.get('precio_venta') or 0
        precio_compra = item_dict.get('precio_compra') or 0
        cantidad = item_dict.get('cantidad') or 0
        if item_dict['tipo'] == 'Producto':
            ganancia_unidad = precio_venta - precio_compra
            item_dict['ganancia_unidad'] = ganancia_unidad
            item_dict['ganancia_total'] = ganancia_unidad * cantidad
            item_dict['por_vencer'] = False
            if item_dict.get('fecha_vencimiento'):
                try:
                    fecha_venc = datetime.strptime(item_dict['fecha_vencimiento'], '%Y-%m-%d').date()
                    if date.today() <= fecha_venc < (date.today() + timedelta(days=90)):
                        item_dict['por_vencer'] = True
                except:
                    pass
        else:
            item_dict['ganancia_unidad'] = 0
            item_dict['ganancia_total'] = 0
            item_dict['por_vencer'] = False
        items.append(item_dict)
    return jsonify(items)

@app.route('/api/citas/horas_ocupadas/<string:fecha>')
@login_required
def api_horas_ocupadas(fecha):
    try:
        datetime.strptime(fecha, '%Y-%m-%d')
    except ValueError:
        return jsonify({'error': 'Formato de fecha inválido'}), 400
    conn = get_db()
    horas_db = conn.execute("SELECT hora_cita FROM citas WHERE fecha_cita = ? AND estado = 'Pendiente'", (fecha,)).fetchall()
    return jsonify([row['hora_cita'] for row in horas_db])

@app.route('/api/citas/historial/buscar')
@login_required
def api_citas_historial_buscar():
    query = request.args.get('q', '', type=str)
    filtro_estado = request.args.get('estado', 'Todos')
    conn = get_db()
    sql = """
        SELECT c.id, c.fecha_cita, c.hora_cita, c.descripcion, c.estado, p.nombre, p.apellido,
               GROUP_CONCAT(i.nombre || ' (' || vd.cantidad || ')', ', ') as productos_usados
        FROM citas c
        JOIN pacientes p ON c.id_paciente = p.id
        LEFT JOIN ventas v ON v.id_cita = c.id AND v.estado != 'Cancelada'
        LEFT JOIN ventas_detalle vd ON vd.id_venta = v.id
        LEFT JOIN inventario i ON vd.id_item_inventario = i.id
        WHERE c.estado IN ('Completada', 'Cancelada')
    """
    params = []
    if query:
        search_term = f"%{query}%"
        sql += " AND (p.nombre LIKE ? OR p.apellido LIKE ?)"
        params.extend([search_term, search_term])
    if filtro_estado in ['Completada', 'Cancelada']:
        sql += " AND c.estado = ?"
        params.append(filtro_estado)
    sql += " GROUP BY c.id ORDER BY c.fecha_cita DESC, c.hora_cita DESC"
    citas_db = conn.execute(sql, tuple(params)).fetchall()
    return jsonify([dict(row) for row in citas_db])

@app.route('/api/inventario/historial/buscar')
@login_required
def api_inventario_historial_buscar():
    query = request.args.get('q', '', type=str)
    conn = get_db()
    search_term = f"%{query}%"
    sql = "SELECT datetime(mi.fecha, 'localtime') as fecha, mi.tipo_movimiento, mi.cantidad, mi.descripcion, i.nombre as nombre_item FROM movimientos_inventario mi JOIN inventario i ON mi.id_item = i.id WHERE i.nombre LIKE ? OR mi.descripcion LIKE ? ORDER BY mi.fecha DESC"
    movimientos_db = conn.execute(sql, (search_term, search_term)).fetchall()
    return jsonify([dict(row) for row in movimientos_db])

@app.route('/api/ventas/buscar')
@login_required
def api_ventas_buscar():
    query = request.args.get('q', '', type=str)
    filtro_fecha = request.args.get('fecha', 'todas')
    conn = get_db()
    sql = """
        SELECT * FROM (
            SELECT
                'Venta' as tipo_registro,
                v.id as id_transaccion,
                datetime(v.fecha_venta, 'localtime') as fecha,
                v.monto_total as monto,
                v.metodo_pago,
                v.referencia_pago,
                v.saldo_pendiente as saldo_pendiente,
                v.estado_pago,
                v.estado as estado_venta,
                p.nombre, p.apellido,
                v.id as id_venta_origen
            FROM ventas v
            JOIN pacientes p ON v.id_paciente = p.id

            UNION ALL

            SELECT
                'Abono' as tipo_registro,
                a.id as id_transaccion,
                datetime(a.fecha_abono, 'localtime') as fecha,
                a.monto as monto,
                a.metodo_pago,
                a.referencia_pago,
                COALESCE(a.saldo_actual_snapshot, 0) as saldo_pendiente,
                v.estado_pago,
                v.estado as estado_venta,
                p.nombre, p.apellido,
                v.id as id_venta_origen
            FROM abonos a
            JOIN ventas v ON a.id_venta = v.id
            JOIN pacientes p ON v.id_paciente = p.id
        ) as historial
        WHERE 1=1
    """
    params = []
    conditions = []
    if query:
        search_term = f"%{query}%"
        conditions.append("(nombre LIKE ? OR apellido LIKE ? OR (nombre || ' ' || apellido) LIKE ?)")
        params.extend([search_term, search_term, search_term])
    if filtro_fecha == 'hoy':
        conditions.append("DATE(fecha) = CURRENT_DATE")
    elif filtro_fecha == 'ayer':
        conditions.append("DATE(fecha) = DATE('now', '-1 day', 'localtime')")
    elif filtro_fecha == 'mes':
        conditions.append("STRFTIME('%Y-%m', fecha) = STRFTIME('%Y-%m', 'now', 'localtime')")
    if conditions:
        sql += f" AND {' AND '.join(conditions)}"
    sql += " ORDER BY fecha DESC"
    items_db = conn.execute(sql, tuple(params)).fetchall()
    return jsonify([dict(row) for row in items_db])

@app.route('/api/pagos/buscar')
@login_required
def api_pagos_buscar():
    query = request.args.get('q', '', type=str)
    filtro_estado = request.args.get('estado', 'todos')
    proximos_vencimientos = request.args.get('proximos', 'false') == 'true'
    db = get_db()
    base_sql = "SELECT * FROM gastos"
    conditions = []
    params = []
    if query:
        conditions.append("(descripcion LIKE ? OR proveedor LIKE ?)")
        params.extend([f"%{query}%", f"%{query}%"])
    if filtro_estado == 'pendiente':
        conditions.append("estado IN ('Pendiente', 'Parcial')")
    elif filtro_estado == 'pagado':
        conditions.append("estado = 'Pagado'")
    if proximos_vencimientos:
        conditions.append("estado IN ('Pendiente', 'Parcial')")
        conditions.append("fecha_vencimiento BETWEEN date('now', 'localtime') AND date('now', '+15 days', 'localtime')")
    if conditions:
        base_sql += " WHERE " + " AND ".join(conditions)
    base_sql += " ORDER BY estado DESC, fecha_vencimiento ASC"
    gastos = db.execute(base_sql, tuple(params)).fetchall()
    return jsonify([dict(row) for row in gastos])

# =============================================================================
# --- 6. RUTAS DE ARCHIVOS (UPLOADS) ---
# =============================================================================

@app.route('/uploads/<path:filename>')
@login_required
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/uploads/productos/<filename>')
@login_required
def uploaded_product_file(filename):
    return send_from_directory(app.config['PRODUCT_UPLOAD_FOLDER'], filename)

@app.route('/uploads/comprobantes/<filename>')
@login_required
def uploaded_expense_file(filename):
    return send_from_directory(app.config['EXPENSE_UPLOAD_FOLDER'], filename)

# =============================================================================
# --- 7. MÓDULOS DE LA APLICACIÓN ---
# =============================================================================

# --- PACIENTES ---
@app.route('/pacientes')
@login_required
def pacientes():
    db = get_db()
    pacientes_db = db.execute('SELECT * FROM pacientes WHERE is_deleted = 0 ORDER BY nombre, apellido').fetchall()
    return render_template('pacientes.html', active_page='pacientes', pacientes=[dict(p, edad=calcular_edad(p['fecha_nacimiento'])) for p in pacientes_db])

@app.route('/pacientes/nuevo', methods=['POST'])
@login_required
def nuevo_paciente():
    n, a = request.form['nombre'], request.form['apellido']
    if not n or not a: flash('El nombre y el apellido son obligatorios.', 'error')
    else:
        db = get_db(); db.execute('INSERT INTO pacientes (nombre, apellido, telefono, fecha_nacimiento) VALUES (?, ?, ?, ?)', (n, a, request.form['telefono'], request.form['fecha_nacimiento'])); db.commit()
        flash('Paciente agregado con éxito.', 'success')
    return redirect(url_for('pacientes'))

@app.route('/paciente/<int:id_paciente>/expedientes')
@login_required
def expedientes_lista(id_paciente):
    db = get_db()
    paciente = db.execute('SELECT * FROM pacientes WHERE id = ?', (id_paciente,)).fetchone()
    expedientes = db.execute('SELECT * FROM expedientes WHERE id_paciente = ? ORDER BY fecha_expediente DESC', (id_paciente,)).fetchall()
    return render_template('expedientes_lista.html', active_page='pacientes', paciente=paciente, expedientes=expedientes)

# --- EXPEDIENTES ---
@app.route('/paciente/<int:id_paciente>/expediente/nuevo', methods=['GET', 'POST'])
@login_required
def expediente_nuevo(id_paciente):
    db = get_db()
    if request.method == 'POST':
        p_nombre = request.form.get('p_nombre')
        p_apellido = request.form.get('p_apellido')
        p_nacimiento = request.form.get('p_fecha_nacimiento')
        p_telefono = request.form.get('p_telefono')
        db.execute("UPDATE pacientes SET nombre=?, apellido=?, fecha_nacimiento=?, telefono=? WHERE id=?",
                   (p_nombre, p_apellido, p_nacimiento, p_telefono, id_paciente))
        cur = db.cursor()
        cur.execute("INSERT INTO expedientes (id_paciente, fecha_expediente, motivo_consulta, enfermedades, vasos_agua, alimentos, problemas_hormonales, alergias, menstruacion, caida_cabello, caspa, horas_sueno, productos_corporales, otros_sintomas, skincare_limpieza, skincare_hidratacion, skincare_ojos, skincare_fotoproteccion, skincare_noche, skincare_tratamiento, sustancias_ingeridas, skincare_nuevo_limpieza, skincare_nuevo_hidratacion, skincare_nuevo_ojos, skincare_nuevo_fotoproteccion, skincare_nuevo_noche, skincare_nuevo_tratamiento) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (id_paciente, request.form['fecha_expediente'], request.form['motivo_consulta'], request.form['enfermedades'], request.form['vasos_agua'], request.form['alimentos'], request.form['problemas_hormonales'], request.form['alergias'], request.form['menstruacion'], request.form['caida_cabello'], request.form['caspa'], request.form['horas_sueno'], request.form['productos_corporales'], request.form['otros_sintomas'], request.form['skincare_limpieza'], request.form['skincare_hidratacion'], request.form['skincare_ojos'], request.form['skincare_fotoproteccion'], request.form['skincare_noche'], request.form['skincare_tratamiento'], request.form['sustancias_ingeridas'], request.form['skincare_nuevo_limpieza'], request.form['skincare_nuevo_hidratacion'], request.form['skincare_nuevo_ojos'], request.form['skincare_nuevo_fotoproteccion'], request.form['skincare_nuevo_noche'], request.form['skincare_nuevo_tratamiento']))
        new_id = cur.lastrowid
        db.commit()
        flash('Expediente creado y datos del paciente actualizados.', 'success')
        return redirect(url_for('expediente_editar', id_expediente=new_id))
    paciente = db.execute('SELECT * FROM pacientes WHERE id = ?', (id_paciente,)).fetchone()
    return render_template('expediente_form.html', active_page='pacientes', accion="Crear", paciente=paciente, expediente=None, sesiones=[], documentos=[], edad=calcular_edad(paciente['fecha_nacimiento']), today=date.today())

@app.route('/expediente/editar/<int:id_expediente>', methods=['GET', 'POST'])
@login_required
def expediente_editar(id_expediente):
    db = get_db()
    exp_actual = db.execute('SELECT id_paciente FROM expedientes WHERE id_expediente = ?', (id_expediente,)).fetchone()
    id_paciente = exp_actual['id_paciente']
    if request.method == 'POST':
        p_nombre = request.form.get('p_nombre')
        p_apellido = request.form.get('p_apellido')
        p_nacimiento = request.form.get('p_fecha_nacimiento')
        p_telefono = request.form.get('p_telefono')
        db.execute("UPDATE pacientes SET nombre=?, apellido=?, fecha_nacimiento=?, telefono=? WHERE id=?",
                   (p_nombre, p_apellido, p_nacimiento, p_telefono, id_paciente))
        db.execute("UPDATE expedientes SET fecha_expediente=?, motivo_consulta=?, enfermedades=?, vasos_agua=?, alimentos=?, problemas_hormonales=?, alergias=?, menstruacion=?, caida_cabello=?, caspa=?, horas_sueno=?, productos_corporales=?, otros_sintomas=?, skincare_limpieza=?, skincare_hidratacion=?, skincare_ojos=?, skincare_fotoproteccion=?, skincare_noche=?, skincare_tratamiento=?, sustancias_ingeridas=?, skincare_nuevo_limpieza=?, skincare_nuevo_hidratacion=?, skincare_nuevo_ojos=?, skincare_nuevo_fotoproteccion=?, skincare_nuevo_noche=?, skincare_nuevo_tratamiento=? WHERE id_expediente=?",
                   (request.form['fecha_expediente'], request.form['motivo_consulta'], request.form['enfermedades'], request.form['vasos_agua'], request.form['alimentos'], request.form['problemas_hormonales'], request.form['alergias'], request.form['menstruacion'], request.form['caida_cabello'], request.form['caspa'], request.form['horas_sueno'], request.form['productos_corporales'], request.form['otros_sintomas'], request.form['skincare_limpieza'], request.form['skincare_hidratacion'], request.form['skincare_ojos'], request.form['skincare_fotoproteccion'], request.form['skincare_noche'], request.form['skincare_tratamiento'], request.form['sustancias_ingeridas'], request.form['skincare_nuevo_limpieza'], request.form['skincare_nuevo_hidratacion'], request.form['skincare_nuevo_ojos'], request.form['skincare_nuevo_fotoproteccion'], request.form['skincare_nuevo_noche'], request.form['skincare_nuevo_tratamiento'], id_expediente))
        db.commit()
        flash('Expediente y datos del paciente actualizados.', 'success')
        return redirect(url_for('expediente_editar', id_expediente=id_expediente))
    expediente = db.execute('SELECT * FROM expedientes WHERE id_expediente = ?', (id_expediente,)).fetchone()
    paciente = db.execute('SELECT * FROM pacientes WHERE id = ?', (expediente['id_paciente'],)).fetchone()
    sesiones = db.execute('SELECT * FROM sesiones WHERE id_expediente = ? ORDER BY fecha_sesion DESC', (id_expediente,)).fetchall()
    documentos = db.execute('SELECT * FROM documentos WHERE id_expediente = ? AND is_deleted = 0 ORDER BY fecha_subida DESC', (id_expediente,)).fetchall()
    return render_template('expediente_form.html', active_page='pacientes', accion="Editar", paciente=paciente, expediente=expediente, sesiones=sesiones, documentos=documentos, edad=calcular_edad(paciente['fecha_nacimiento']), today=date.today())

@app.route('/sesion/nueva/<int:id_expediente>', methods=['POST'])
@login_required
def sesion_nueva(id_expediente):
    db = get_db()
    db.execute('INSERT INTO sesiones (id_expediente, fecha_sesion, tratamiento, resultados) VALUES (?, ?, ?, ?)', (id_expediente, request.form['fecha_sesion'], request.form['tratamiento'], request.form['resultados']))
    db.commit()
    flash('Nueva sesión agregada.', 'success')
    return redirect(url_for('expediente_editar', id_expediente=id_expediente) + '#documentos')

@app.route('/documento/subir/<int:id_expediente>', methods=['POST'])
@login_required
def subir_documento(id_expediente):
    files = request.files.getlist('documento')
    if not files or files[0].filename == '':
        flash('No se seleccionó ningún archivo.', 'error')
        return redirect(url_for('expediente_editar', id_expediente=id_expediente) + '#documentos')
    db = get_db()
    uploads = 0
    for file in files:
        if file and allowed_file(file.filename):
            unique_name = save_unique_file(file, app.config['UPLOAD_FOLDER'])
            db.execute('INSERT INTO documentos (id_expediente, nombre_archivo, ruta_archivo) VALUES (?, ?, ?)', (id_expediente, unique_name, unique_name))
            uploads += 1
    if uploads > 0:
        db.commit()
        flash(f'{uploads} documento(s) subido(s) con éxito.', 'success')
    return redirect(url_for('expediente_editar', id_expediente=id_expediente) + '#documentos')

# --- INVENTARIO ---
@app.route('/inventario')
@login_required
def inventario():
    conn = get_db()
    items_db = conn.execute('SELECT * FROM inventario WHERE is_deleted = 0 ORDER BY tipo DESC, nombre').fetchall()
    items = []
    for item in items_db:
        item_dict = dict(item)
        precio_venta, precio_compra, cantidad = item_dict.get('precio_venta') or 0, item_dict.get('precio_compra') or 0, item_dict.get('cantidad') or 0
        if item_dict['tipo'] == 'Producto':
            ganancia_unidad = precio_venta - precio_compra
            item_dict['ganancia_unidad'] = ganancia_unidad
            item_dict['ganancia_total'] = ganancia_unidad * cantidad
            item_dict['por_vencer'] = False
            if item_dict.get('fecha_vencimiento'):
                try:
                    fecha_venc = datetime.strptime(item_dict['fecha_vencimiento'], '%Y-%m-%d').date()
                    if date.today() <= fecha_venc < (date.today() + timedelta(days=20)):
                        item_dict['por_vencer'] = True
                except:
                    pass
        else:
            item_dict['ganancia_unidad'] = 0
            item_dict['ganancia_total'] = 0
            item_dict['por_vencer'] = False
        items.append(item_dict)
    return render_template('inventario.html', active_page='inventario', items=items)

@app.route('/inventario/exportar-xlsx')
@login_required
def inventario_exportar_xlsx():
    conn = get_db()
    productos = conn.execute(
        "SELECT nombre, cantidad, precio_compra, precio_venta, fecha_vencimiento "
        "FROM inventario WHERE is_deleted = 0 AND tipo = 'Producto' AND COALESCE(cantidad, 0) > 0 "
        "ORDER BY nombre"
    ).fetchall()

    COLOR_PETROLEO = "1F7A7A"
    COLOR_ARENA = "EDC988"
    COLOR_FONDO = "F3E6E0"

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = ["Producto", "Cantidad", "Costo Unit. (Q)", "Venta Unit. (Q)", "Costo Total (Q)", "Venta Total (Q)", "Vencimiento"]
    ws.merge_cells('A1:G1')
    titulo = ws['A1']
    titulo.value = "REPORTE GENERAL DE INVENTARIO"
    titulo.font = Font(name="Calibri", size=16, bold=True, color=COLOR_PETROLEO)
    titulo.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:G2')
    subtitulo = ws['A2']
    subtitulo.value = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    subtitulo.font = Font(name="Calibri", size=10, italic=True, color="666666")
    subtitulo.alignment = Alignment(horizontal='center')

    header_row = 4
    thin = Side(style='thin', color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COLOR_PETROLEO, end_color=COLOR_PETROLEO, fill_type="solid")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    total_costo = 0.0
    total_venta = 0.0
    row_idx = header_row + 1
    for prod in productos:
        cantidad = prod['cantidad'] or 0
        costo_unit = prod['precio_compra'] or 0
        venta_unit = prod['precio_venta'] or 0
        costo_total = costo_unit * cantidad
        venta_total = venta_unit * cantidad
        total_costo += costo_total
        total_venta += venta_total

        valores = [prod['nombre'], cantidad, costo_unit, venta_unit, costo_total, venta_total, prod['fecha_vencimiento'] or 'N/A']
        for col_idx, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = border
            if col_idx in (3, 4, 5, 6):
                cell.number_format = '"Q"#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal='center')
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal='center')
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color=COLOR_FONDO, end_color=COLOR_FONDO, fill_type="solid")
        row_idx += 1

    total_row = row_idx
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    total_label = ws.cell(row=total_row, column=1, value="TOTALES")
    total_label.font = Font(bold=True, color=COLOR_PETROLEO)
    total_label.alignment = Alignment(horizontal='right', vertical='center')
    total_label.fill = PatternFill(start_color=COLOR_ARENA, end_color=COLOR_ARENA, fill_type="solid")
    for col in range(1, 5):
        ws.cell(row=total_row, column=col).fill = PatternFill(start_color=COLOR_ARENA, end_color=COLOR_ARENA, fill_type="solid")

    for col_idx, total_val in ((5, total_costo), (6, total_venta)):
        cell = ws.cell(row=total_row, column=col_idx, value=total_val)
        cell.number_format = '"Q"#,##0.00'
        cell.font = Font(bold=True, color=COLOR_PETROLEO)
        cell.fill = PatternFill(start_color=COLOR_ARENA, end_color=COLOR_ARENA, fill_type="solid")
        cell.alignment = Alignment(horizontal='right')
    ws.cell(row=total_row, column=7).fill = PatternFill(start_color=COLOR_ARENA, end_color=COLOR_ARENA, fill_type="solid")

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:G{header_row}"

    column_widths = {1: 34, 2: 11, 3: 15, 4: 15, 5: 16, 6: 16, 7: 14}
    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"Inventario_{date.today().isoformat()}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/inventario/item/nuevo', methods=['GET', 'POST'])
@login_required
def inventario_nuevo():
    if request.method == 'POST':
        conn = get_db()
        procesar_formulario_inventario(conn)
        flash('Ítem de inventario creado con éxito.', 'success')
        return redirect(url_for('inventario'))
    return render_template('inventario_form.html', active_page='inventario', accion="Añadir", item=None)

@app.route('/inventario/item/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def inventario_editar(id):
    conn = get_db()
    if request.method == 'POST':
        procesar_formulario_inventario(conn, item_id=id)
        flash('Ítem de inventario actualizado con éxito.', 'success')
        return redirect(url_for('inventario'))
    item = conn.execute('SELECT * FROM inventario WHERE id = ?', (id,)).fetchone()
    return render_template('inventario_form.html', active_page='inventario', accion="Editar", item=item)

@app.route('/inventario/stock/agregar/<int:id>', methods=['POST'])
@login_required
def inventario_agregar_stock(id):
    cantidad, desc = request.form.get('cantidad_a_agregar', 0, type=int), request.form.get('descripcion_movimiento', 'Entrada de stock manual.')
    if cantidad > 0:
        conn = get_db()
        # CORRECCIÓN: Usar COALESCE y CAST para asegurar que el incremento sea numérico
        conn.execute('UPDATE inventario SET cantidad = CAST(COALESCE(cantidad, 0) AS INTEGER) + ? WHERE id = ?', (cantidad, id))
        conn.execute("INSERT INTO movimientos_inventario (id_item, tipo_movimiento, cantidad, descripcion) VALUES (?, 'Entrada', ?, ?)", (id, cantidad, desc))
        conn.commit()
        flash(f'Se agregaron {cantidad} unidades al stock.', 'success')
    else:
        flash('La cantidad a agregar debe ser mayor que cero.', 'error')
    return redirect(url_for('inventario'))

@app.route('/inventario/stock/remover/<int:id>', methods=['POST'])
@login_required
@admin_required
def inventario_remover_stock(id):
    cantidad_a_remover = request.form.get('cantidad_a_remover', 0, type=int)
    descripcion = request.form.get('descripcion_movimiento', 'Salida de stock manual.')
    conn = get_db()
    item_actual = conn.execute('SELECT cantidad FROM inventario WHERE id = ?', (id,)).fetchone()
    if not item_actual or (item_actual['cantidad'] or 0) < cantidad_a_remover:
        flash('No hay suficiente stock para remover la cantidad indicada.', 'error')
        return redirect(url_for('inventario'))
    if cantidad_a_remover > 0:
        # CORRECCIÓN: Usar COALESCE y CAST para asegurar que el decremento sea numérico
        conn.execute('UPDATE inventario SET cantidad = CAST(COALESCE(cantidad, 0) AS INTEGER) - ? WHERE id = ?', (cantidad_a_remover, id))
        conn.execute("INSERT INTO movimientos_inventario (id_item, tipo_movimiento, cantidad, descripcion) VALUES (?, 'Salida', ?, ?)", (id, cantidad_a_remover, descripcion))
        conn.commit()
        flash(f'Se removieron {cantidad_a_remover} unidades del stock.', 'success')
    else:
        flash('La cantidad a remover debe ser mayor que cero.', 'error')
    return redirect(url_for('inventario'))

@app.route('/inventario/historial')
@login_required
def inventario_historial():
    return render_template('inventario_historial.html', active_page='inventario_historial')

# --- CITAS ---
@app.route('/citas')
@login_required
def citas():
    conn = get_db()
    citas_pendientes = conn.execute("SELECT c.id, c.fecha_cita, c.hora_cita, c.descripcion, c.estado, p.nombre, p.apellido FROM citas c JOIN pacientes p ON c.id_paciente = p.id WHERE c.estado = 'Pendiente' ORDER BY c.fecha_cita ASC, c.hora_cita ASC").fetchall()
    return render_template('citas.html', active_page='citas', citas_pendientes=citas_pendientes)

@app.route('/citas/historial')
@login_required
def citas_historial():
    return render_template('citas_historial.html', active_page='citas')

@app.route('/citas/nueva', methods=['GET', 'POST'])
@login_required
def cita_nueva():
    conn = get_db()
    if request.method == 'POST':
        id_paciente, fecha, hora, desc = request.form.get('id_paciente'), request.form.get('fecha_cita'), request.form.get('hora_cita'), request.form.get('descripcion')
        if not id_paciente or not fecha or not hora:
            flash('Paciente, fecha y hora son obligatorios.', 'error')
        else:
            cita_existente = conn.execute('SELECT id FROM citas WHERE fecha_cita = ? AND hora_cita = ? AND estado = "Pendiente"', (fecha, hora)).fetchone()
            if cita_existente:
                flash(f'Ya existe una cita agendada para el {fecha} a las {hora}.', 'error')
            else:
                conn.execute('INSERT INTO citas (id_paciente, fecha_cita, hora_cita, descripcion) VALUES (?, ?, ?, ?)', (id_paciente, fecha, hora, desc))
                conn.commit()
                flash('Cita agendada con éxito.', 'success')
                return redirect(url_for('citas'))
    pacientes = conn.execute('SELECT id, nombre, apellido FROM pacientes WHERE is_deleted = 0 ORDER BY nombre, apellido').fetchall()
    return render_template('cita_form.html', active_page='citas', accion="Agendar", pacientes=pacientes, cita=request.form)

# --- EDITAR CITA ---
@app.route('/citas/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def cita_editar(id):
    conn = get_db()
    if request.method == 'POST':
        id_paciente, fecha, hora, desc = request.form.get('id_paciente'), request.form.get('fecha_cita'), request.form.get('hora_cita'), request.form.get('descripcion')
        if not id_paciente or not fecha or not hora:
            flash('Todos los campos son obligatorios.', 'error')
        else:
            conn.execute('UPDATE citas SET id_paciente=?, fecha_cita=?, hora_cita=?, descripcion=? WHERE id=?', (id_paciente, fecha, hora, desc, id))
            conn.commit()
            flash('Cita actualizada correctamente.', 'success')
            return redirect(url_for('citas'))

    cita = conn.execute("""
        SELECT c.*, p.nombre, p.apellido
        FROM citas c
        JOIN pacientes p ON c.id_paciente = p.id
        WHERE c.id = ?
    """, (id,)).fetchone()

    pacientes = conn.execute('SELECT id, nombre, apellido FROM pacientes WHERE is_deleted = 0 ORDER BY nombre, apellido').fetchall()
    return render_template('cita_form.html', active_page='citas', accion="Editar", pacientes=pacientes, cita=cita)

@app.route('/citas/atender/<int:id_cita>')
@login_required
def atender_cita(id_cita):
    db = get_db()
    caja_abierta = get_caja_abierta(db) # <--- AGREGADO
    cita = db.execute("SELECT c.*, p.nombre, p.apellido FROM citas c JOIN pacientes p ON c.id_paciente = p.id WHERE c.id = ?", (id_cita,)).fetchone()
    if not cita:
        flash('Cita no encontrada.', 'error')
        return redirect(url_for('citas'))
    return render_template('venta_form.html', active_page='citas', accion="Atender", cita_data=cita, caja_abierta=caja_abierta) # <--- ACTUALIZADO

@app.route('/citas/completar/<int:id>', methods=['POST'])
@login_required
def cita_completar(id):
    conn = get_db()
    conn.execute("UPDATE citas SET estado = 'Completada' WHERE id = ?", (id,))
    conn.commit()
    flash('Cita marcada como completada.', 'success')
    return redirect(url_for('citas'))

@app.route('/citas/cancelar/<int:id>', methods=['POST'])
@login_required
def cita_cancelar(id):
    conn = get_db()
    conn.execute("UPDATE citas SET estado = 'Cancelada' WHERE id = ?", (id,))
    conn.commit()
    flash('Cita cancelada.', 'success')
    return redirect(url_for('citas'))

# --- VENTAS, CRÉDITOS Y CAJA ---
@app.route('/ventas')
@login_required
def ventas():
    db = get_db()
    ventas = db.execute("""
        SELECT * FROM (
            SELECT 'Venta' as tipo_registro, v.id as id_transaccion, datetime(v.fecha_venta, 'localtime') as fecha, v.monto_total as monto, v.metodo_pago, v.referencia_pago, v.saldo_pendiente, v.estado_pago, v.estado as estado_venta, p.nombre, p.apellido, v.id as id_venta_origen
            FROM ventas v JOIN pacientes p ON v.id_paciente = p.id
            UNION ALL
            SELECT 'Abono' as tipo_registro, a.id as id_transaccion, datetime(a.fecha_abono, 'localtime') as fecha, a.monto as monto, a.metodo_pago, a.referencia_pago, COALESCE(a.saldo_actual_snapshot, 0) as saldo_pendiente, v.estado_pago, v.estado as estado_venta, p.nombre, p.apellido, v.id as id_venta_origen
            FROM abonos a JOIN ventas v ON a.id_venta = v.id JOIN pacientes p ON v.id_paciente = p.id
        ) ORDER BY fecha DESC
    """).fetchall()
    caja_abierta = get_caja_abierta(db)
    return render_template('ventas.html', active_page='ventas', ventas=ventas, caja=caja_abierta)

@app.route('/creditos')
@login_required
def creditos():
    db = get_db()
    creditos = db.execute("""
        SELECT v.id, datetime(v.fecha_venta, 'localtime') as fecha_venta, v.monto_total, v.saldo_pendiente, v.estado_pago, p.nombre, p.apellido, p.telefono
        FROM ventas v JOIN pacientes p ON v.id_paciente = p.id
        WHERE v.metodo_pago = 'Crédito' AND v.estado != 'Cancelada' ORDER BY v.estado_pago DESC, v.fecha_venta DESC
    """).fetchall()
    total_por_cobrar = sum(c['saldo_pendiente'] for c in creditos)
    return render_template('creditos.html', active_page='creditos', creditos=creditos, total_por_cobrar=total_por_cobrar)

@app.route('/caja/abrir', methods=['POST'])
@login_required
def abrir_caja():
    monto_inicial = request.form.get('monto_inicial', 0, type=float)
    db = get_db()
    if get_caja_abierta(db):
        flash('Ya existe una caja abierta.', 'error')
        return redirect(url_for('ventas'))
    db.execute("INSERT INTO cierres_caja (monto_inicial, estado) VALUES (?, 'Abierta')", (monto_inicial,))
    db.commit()
    flash('Caja abierta correctamente.', 'success')
    return redirect(url_for('ventas'))

@app.route('/caja/pre-cierre')
@login_required
def pre_cierre_caja():
    db = get_db()
    caja = get_caja_abierta(db)
    if not caja: return jsonify({'error': 'No hay caja abierta'}), 400
    fecha_apertura = caja['fecha_apertura']
    ventas_resumen = db.execute("SELECT SUM(CASE WHEN metodo_pago = 'Efectivo' THEN monto_total ELSE 0 END) as efectivo, SUM(CASE WHEN metodo_pago NOT IN ('Efectivo', 'Crédito') THEN monto_total ELSE 0 END) as otros FROM ventas WHERE fecha_venta >= ? AND metodo_pago != 'Crédito' AND estado != 'Cancelada'", (fecha_apertura,)).fetchone()
    abonos_resumen = db.execute("SELECT SUM(CASE WHEN a.metodo_pago = 'Efectivo' THEN a.monto ELSE 0 END) as efectivo, SUM(CASE WHEN a.metodo_pago != 'Efectivo' THEN a.monto ELSE 0 END) as otros FROM abonos a JOIN ventas v ON a.id_venta = v.id WHERE a.fecha_abono >= ? AND v.estado != 'Cancelada'", (fecha_apertura,)).fetchone()
    v_efectivo = (ventas_resumen['efectivo'] or 0) + (abonos_resumen['efectivo'] or 0)
    v_otros = (ventas_resumen['otros'] or 0) + (abonos_resumen['otros'] or 0)
    total_calculado = caja['monto_inicial'] + v_efectivo
    return jsonify({'monto_inicial': caja['monto_inicial'], 'ventas_efectivo': v_efectivo, 'ventas_otros': v_otros, 'total_calculado': total_calculado})

@app.route('/caja/cerrar', methods=['POST'])
@login_required
def cerrar_caja():
    db = get_db()
    caja = get_caja_abierta(db)
    if not caja:
        flash('No hay una caja abierta.', 'error')
        return redirect(url_for('ventas'))
    monto_final_real = request.form.get('monto_final_real', 0, type=float)
    fecha_apertura = caja['fecha_apertura']
    ventas_resumen = db.execute("SELECT SUM(CASE WHEN metodo_pago = 'Efectivo' THEN monto_total ELSE 0 END) as efectivo, SUM(CASE WHEN metodo_pago NOT IN ('Efectivo', 'Crédito') THEN monto_total ELSE 0 END) as otros FROM ventas WHERE fecha_venta >= ? AND metodo_pago != 'Crédito' AND estado != 'Cancelada'", (fecha_apertura,)).fetchone()
    abonos_resumen = db.execute("SELECT SUM(CASE WHEN a.metodo_pago = 'Efectivo' THEN a.monto ELSE 0 END) as efectivo, SUM(CASE WHEN a.metodo_pago != 'Efectivo' THEN a.monto ELSE 0 END) as otros FROM abonos a JOIN ventas v ON a.id_venta = v.id WHERE a.fecha_abono >= ? AND v.estado != 'Cancelada'", (fecha_apertura,)).fetchone()
    v_efectivo = (ventas_resumen['efectivo'] or 0) + (abonos_resumen['efectivo'] or 0)
    v_otros = (ventas_resumen['otros'] or 0) + (abonos_resumen['otros'] or 0)
    monto_final_calculado = caja['monto_inicial'] + v_efectivo
    diferencia = monto_final_real - monto_final_calculado
    db.execute("UPDATE cierres_caja SET fecha_cierre = CURRENT_TIMESTAMP, total_ventas_efectivo = ?, total_ventas_otros = ?, monto_final_calculado = ?, monto_final_real = ?, diferencia = ?, estado = 'Cerrada' WHERE id = ?", (v_efectivo, v_otros, monto_final_calculado, monto_final_real, diferencia, caja['id']))
    db.commit()
    flash('Caja cerrada correctamente.', 'success')
    return redirect(url_for('ventas'))

@app.route('/ventas/cancelar/<int:id_venta>', methods=['POST'])
@login_required
def cancelar_venta(id_venta):
    db = get_db()
    venta = db.execute('SELECT * FROM ventas WHERE id = ?', (id_venta,)).fetchone()
    if not venta:
        flash('Venta no encontrada.', 'error'); return redirect(url_for('ventas'))
    if venta['estado'] == 'Cancelada':
        flash('Esta venta ya está cancelada.', 'warning'); return redirect(url_for('venta_detalle', id_venta=id_venta))
    abonos_count = db.execute("SELECT COUNT(*) FROM abonos WHERE id_venta = ?", (id_venta,)).fetchone()[0]
    if abonos_count > 0:
        flash('No se puede cancelar una venta que ya tiene abonos registrados.', 'error'); return redirect(url_for('venta_detalle', id_venta=id_venta))
    try:
        detalles = db.execute("SELECT * FROM ventas_detalle WHERE id_venta = ?", (id_venta,)).fetchall()
        for item in detalles:
            prod = db.execute("SELECT tipo FROM inventario WHERE id = ?", (item['id_item_inventario'],)).fetchone()
            # CORRECCIÓN: Manejo de tipos para cancelación robusto
            tipo_item = str(prod['tipo']).strip().lower() if (prod and prod['tipo']) else ""
            if tipo_item == 'producto':
                db.execute("UPDATE inventario SET cantidad = CAST(COALESCE(cantidad, 0) AS INTEGER) + ? WHERE id = ?", (item['cantidad'], item['id_item_inventario']))
                db.execute("INSERT INTO movimientos_inventario (id_item, tipo_movimiento, cantidad, descripcion) VALUES (?, 'Entrada por Cancelación', ?, ?)", (item['id_item_inventario'], item['cantidad'], f'Cancelación Venta #{id_venta}'))
        db.execute("UPDATE ventas SET estado = 'Cancelada', saldo_pendiente = 0 WHERE id = ?", (id_venta,))
        db.commit()
        flash(f'Venta #{id_venta} cancelada correctamente. Stock devuelto.', 'success')
    except Exception as e:
        db.rollback(); flash(f'Error al cancelar venta: {e}', 'error')
    return redirect(url_for('venta_detalle', id_venta=id_venta))

@app.route('/ventas/nueva', methods=['GET', 'POST'])
@login_required
def venta_nueva():
    db = get_db()
    caja_abierta = get_caja_abierta(db) # <--- AGREGADO

    if request.method == 'POST':
        id_paciente = request.form.get('id_paciente')
        item_ids = request.form.getlist('item_ids[]')
        cantidades = request.form.getlist('quantities[]')
        descuentos_items = request.form.getlist('discounts[]')
        descuento_general = request.form.get('descuento_general', 0, type=float)
        metodo_pago = request.form.get('metodo_pago')
        referencia_pago = request.form.get('referencia_pago', '')
        id_cita = request.form.get('id_cita')
        if not id_paciente or not item_ids or not metodo_pago:
            flash('Datos incompletos.', 'error'); return redirect(url_for('venta_nueva'))
        if metodo_pago != 'Efectivo' and metodo_pago != 'Crédito' and not referencia_pago:
            flash('Referencia obligatoria.', 'error'); return redirect(url_for('venta_nueva'))
        try:
            subtotal = 0; items_para_venta = []
            for i, item_id in enumerate(item_ids):
                cantidad = int(cantidades[i]); descuento = float(descuentos_items[i]) if descuentos_items[i] else 0.0
                item = db.execute('SELECT * FROM inventario WHERE id = ?', (item_id,)).fetchone()

                # Normalización de tipo
                tipo_item = str(item['tipo']).strip().lower() if item['tipo'] else ""

                # Verificar stock si es producto
                if tipo_item == 'producto' and (item['cantidad'] or 0) < cantidad:
                    flash(f"No hay stock suficiente para '{item['nombre']}'.", 'error'); return redirect(url_for('venta_nueva'))

                precio_unitario = item['precio_venta']
                subtotal += (precio_unitario - descuento) * cantidad
                items_para_venta.append({'item': item, 'cantidad': cantidad, 'precio_unitario': precio_unitario, 'descuento': descuento, 'tipo_normalizado': tipo_item})

            monto_total = subtotal - descuento_general
            saldo_pendiente = monto_total if metodo_pago == 'Crédito' else 0
            estado_pago = 'Pendiente' if metodo_pago == 'Crédito' else 'Pagado'

            cur = db.cursor()
            cur.execute("INSERT INTO ventas (id_paciente, monto_total, descuento_total, metodo_pago, referencia_pago, saldo_pendiente, estado_pago, id_cita) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (id_paciente, monto_total, descuento_general, metodo_pago, referencia_pago, saldo_pendiente, estado_pago, id_cita))
            id_venta = cur.lastrowid

            for venta_item in items_para_venta:
                item, cantidad, descuento = venta_item['item'], venta_item['cantidad'], venta_item['descuento']
                cur.execute("INSERT INTO ventas_detalle (id_venta, id_item_inventario, cantidad, precio_unitario, descuento) VALUES (?, ?, ?, ?, ?)", (id_venta, item['id'], cantidad, venta_item['precio_unitario'], descuento))

                # CORRECCIÓN: Restar inventario con casting a INTEGER y COALESCE para evitar errores de tipo o nulos
                if venta_item['tipo_normalizado'] == 'producto':
                    cur.execute("""
                        UPDATE inventario
                        SET cantidad = CAST(COALESCE(cantidad, 0) AS INTEGER) - ?
                        WHERE id = ?
                    """, (cantidad, item['id']))
                    cur.execute("INSERT INTO movimientos_inventario (id_item, tipo_movimiento, cantidad, descripcion) VALUES (?, 'Salida por Venta', ?, ?)", (item['id'], cantidad, f'Venta #{id_venta}'))

            if id_cita: cur.execute("UPDATE citas SET estado = 'Completada' WHERE id = ?", (id_cita,))
            db.commit(); flash(f'Venta #{id_venta} registrada.', 'success'); return redirect(url_for('ventas'))
        except Exception as e: db.rollback(); flash(f'Error: {e}', 'error'); return redirect(url_for('venta_nueva'))
    return render_template('venta_form.html', active_page='ventas', accion="Registrar", caja_abierta=caja_abierta) # <--- ACTUALIZADO

@app.route('/ventas/abonar/<int:id_venta>', methods=['POST'])
@login_required
def abonar_venta(id_venta):
    monto = request.form.get('monto_abono', type=float)
    metodo = request.form.get('metodo_abono'); referencia = request.form.get('referencia_abono', '')
    origen = request.form.get('origen_url', 'ventas')
    if not monto or monto <= 0: flash('Monto inválido.', 'error'); return redirect(url_for('ventas'))
    db = get_db(); venta = db.execute('SELECT * FROM ventas WHERE id = ?', (id_venta,)).fetchone()
    if not venta or monto > venta['saldo_pendiente']: flash('Venta no encontrada o monto excesivo.', 'error'); return redirect(url_for('ventas'))
    nuevo_saldo = venta['saldo_pendiente'] - monto
    nuevo_estado = 'Pagado' if nuevo_saldo <= 0 else 'Pendiente'
    try:
        db.execute('INSERT INTO abonos (id_venta, monto, metodo_pago, referencia_pago, saldo_actual_snapshot) VALUES (?, ?, ?, ?, ?)', (id_venta, monto, metodo, referencia, nuevo_saldo))
        db.execute('UPDATE ventas SET saldo_pendiente = ?, estado_pago = ? WHERE id = ?', (nuevo_saldo, nuevo_estado, id_venta))
        db.commit(); flash('Abono registrado.', 'success')
    except Exception as e: db.rollback(); flash(f'Error: {e}', 'error')
    return redirect(url_for(origen))

@app.route('/venta/detalle/<int:id_venta>')
@login_required
def venta_detalle(id_venta):
    db = get_db()
    venta = db.execute("SELECT v.id, datetime(v.fecha_venta, 'localtime') as fecha_venta, v.monto_total, v.descuento_total, v.metodo_pago, v.referencia_pago, v.saldo_pendiente, v.estado_pago, v.estado, p.nombre, p.apellido, p.telefono FROM ventas v JOIN pacientes p ON v.id_paciente = p.id WHERE v.id = ?", (id_venta,)).fetchone()
    if not venta: flash('No encontrada.', 'error'); return redirect(url_for('ventas'))
    detalles = db.execute("SELECT vd.*, i.nombre as nombre_item FROM ventas_detalle vd JOIN inventario i ON vd.id_item_inventario = i.id WHERE vd.id_venta = ?", (id_venta,)).fetchall()
    abonos = db.execute("SELECT a.id, datetime(a.fecha_abono, 'localtime') as fecha_abono, a.monto, a.metodo_pago, a.referencia_pago, a.saldo_actual_snapshot FROM abonos a WHERE a.id_venta = ? ORDER BY a.fecha_abono DESC", (id_venta,)).fetchall()
    subtotal = sum((d['precio_unitario'] - d['descuento']) * d['cantidad'] for d in detalles)
    return render_template('venta_detalle.html', active_page='ventas', venta=venta, detalles=detalles, abonos=abonos, subtotal=subtotal)

# --- REPORTES ---
@app.route('/caja/historial')
@login_required
def historial_cierres():
    db = get_db()
    cierres = db.execute("SELECT id, datetime(fecha_apertura, 'localtime') as fecha_apertura, datetime(fecha_cierre, 'localtime') as fecha_cierre, monto_inicial, monto_final_calculado, monto_final_real, diferencia, estado FROM cierres_caja ORDER BY fecha_apertura DESC").fetchall()
    return render_template('cierres_historial.html', active_page='reportes', cierres=cierres)

@app.route('/reporte/financiero', methods=['POST'])
@login_required
def generar_reporte_financiero():
    fecha_inicio = request.form.get('fecha_inicio')
    fecha_fin = request.form.get('fecha_fin')
    db = get_db()
    ventas_data = db.execute("SELECT COUNT(id) as total_trx, SUM(monto_total) as total_vendido FROM ventas WHERE DATE(fecha_venta) BETWEEN ? AND ? AND estado != 'Cancelada'", (fecha_inicio, fecha_fin)).fetchone()
    ganancia_data = db.execute("SELECT SUM((vd.precio_unitario - vd.descuento) * vd.cantidad) - SUM(i.precio_compra * vd.cantidad) as ganancia_bruta FROM ventas_detalle vd JOIN ventas v ON vd.id_venta = v.id JOIN inventario i ON vd.id_item_inventario = i.id WHERE DATE(v.fecha_venta) BETWEEN ? AND ? AND v.estado != 'Cancelada'", (fecha_inicio, fecha_fin)).fetchone()
    ingresos_data = db.execute("SELECT SUM(CASE WHEN metodo_pago = 'Efectivo' THEN monto ELSE 0 END) as total_efectivo, SUM(CASE WHEN metodo_pago NOT IN ('Efectivo', 'Crédito') THEN monto ELSE 0 END) as total_bancos FROM (SELECT metodo_pago, monto_total as monto FROM ventas WHERE DATE(fecha_venta) BETWEEN ? AND ? AND metodo_pago != 'Crédito' AND estado != 'Cancelada' UNION ALL SELECT metodo_pago, monto FROM abonos WHERE DATE(fecha_abono) BETWEEN ? AND ?) ", (fecha_inicio, fecha_fin, fecha_inicio, fecha_fin)).fetchone()
    creditos_data = db.execute("SELECT SUM(monto_total) as total_credito FROM ventas WHERE DATE(fecha_venta) BETWEEN ? AND ? AND metodo_pago = 'Crédito' AND estado != 'Cancelada'", (fecha_inicio, fecha_fin)).fetchone()
    tipo_data = db.execute("SELECT SUM(CASE WHEN i.tipo = 'Producto' THEN (vd.precio_unitario - vd.descuento) * vd.cantidad ELSE 0 END) as total_productos, SUM(CASE WHEN i.tipo = 'Servicio' THEN (vd.precio_unitario - vd.descuento) * vd.cantidad ELSE 0 END) as total_servicios FROM ventas_detalle vd JOIN ventas v ON vd.id_venta = v.id JOIN inventario i ON vd.id_item_inventario = i.id WHERE DATE(v.fecha_venta) BETWEEN ? AND ? AND v.estado != 'Cancelada'", (fecha_inicio, fecha_fin)).fetchone()

    reporte = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'total_ventas': ventas_data['total_vendido'] or 0, 'trx_count': ventas_data['total_trx'] or 0, 'ganancia_bruta': ganancia_data['ganancia_bruta'] or 0, 'ingreso_efectivo': ingresos_data['total_efectivo'] or 0, 'ingreso_bancos': ingresos_data['total_bancos'] or 0, 'creditos_nuevos': creditos_data['total_credito'] or 0, 'total_productos': tipo_data['total_productos'] or 0, 'total_servicios': tipo_data['total_servicios'] or 0}
    sql_base = "SELECT v.id, datetime(v.fecha_venta, 'localtime') as fecha, p.nombre, p.apellido, v.monto_total, v.metodo_pago, v.referencia_pago, v.saldo_pendiente, GROUP_CONCAT(vd.cantidad || 'x ' || i.nombre, ', ') as detalles_productos FROM ventas v JOIN pacientes p ON v.id_paciente = p.id JOIN ventas_detalle vd ON v.id = vd.id_venta JOIN inventario i ON vd.id_item_inventario = i.id WHERE DATE(v.fecha_venta) BETWEEN ? AND ? AND v.estado != 'Cancelada'"
    lista_efectivo = db.execute(sql_base + " AND v.metodo_pago = 'Efectivo' GROUP BY v.id ORDER BY v.fecha_venta", (fecha_inicio, fecha_fin)).fetchall()
    lista_otros = db.execute(sql_base + " AND v.metodo_pago NOT IN ('Efectivo', 'Crédito') GROUP BY v.id ORDER BY v.fecha_venta", (fecha_inicio, fecha_fin)).fetchall()
    lista_credito = db.execute(sql_base + " AND v.metodo_pago = 'Crédito' GROUP BY v.id ORDER BY v.fecha_venta", (fecha_inicio, fecha_fin)).fetchall()

    return render_template('reporte_financiero.html', reporte=reporte, lista_efectivo=lista_efectivo, lista_otros=lista_otros, lista_credito=lista_credito)

# --- GASTOS ---
@app.route('/pagos')
@login_required
def pagos():
    db = get_db(); gastos = db.execute("SELECT * FROM gastos ORDER BY estado DESC, fecha_vencimiento ASC").fetchall()
    total_pendiente = db.execute("SELECT SUM(saldo_pendiente) FROM gastos WHERE estado != 'Pagado'").fetchone()[0] or 0
    return render_template('pagos.html', active_page='pagos', gastos=gastos, total_pendiente=total_pendiente)

@app.route('/pagos/nuevo', methods=['POST'])
@login_required
def nuevo_gasto():
    d, p, m, fe, fv = request.form.get('descripcion'), request.form.get('proveedor'), request.form.get('monto', type=float), request.form.get('fecha_emision'), request.form.get('fecha_vencimiento')
    if not d or not m: flash('Descripción y monto son obligatorios.', 'error'); return redirect(url_for('pagos'))
    db = get_db(); db.execute("INSERT INTO gastos (descripcion, proveedor, monto, fecha_emision, fecha_vencimiento, saldo_pendiente) VALUES (?, ?, ?, ?, ?, ?)", (d, p, m, fe, fv, m)); db.commit()
    flash('Gasto registrado.', 'success'); return redirect(url_for('pagos'))

@app.route('/pagos/pagar/<int:id_gasto>', methods=['POST'])
@login_required
def pagar_gasto(id_gasto):
    fecha_pago = request.form.get('fecha_pago', date.today().strftime('%Y-%m-%d'))
    archivo = request.files.get('comprobante'); filename = None
    if archivo and archivo.filename != '':
        filename = save_unique_file(archivo, app.config['EXPENSE_UPLOAD_FOLDER'])
    db = get_db()
    if filename: db.execute("UPDATE gastos SET estado = 'Pagado', fecha_pago = ?, comprobante_ruta = ?, saldo_pendiente = 0 WHERE id = ?", (fecha_pago, filename, id_gasto))
    else: db.execute("UPDATE gastos SET estado = 'Pagado', fecha_pago = ?, saldo_pendiente = 0 WHERE id = ?", (fecha_pago, id_gasto))
    db.commit(); flash('Gasto pagado.', 'success'); return redirect(url_for('pagos'))

@app.route('/pagos/abonar/<int:id_gasto>', methods=['POST'])
@login_required
def abonar_gasto(id_gasto):
    monto = request.form.get('monto_abono', type=float)
    metodo = request.form.get('metodo_abono'); referencia = request.form.get('referencia_abono', '')
    db = get_db(); gasto = db.execute('SELECT * FROM gastos WHERE id = ?', (id_gasto,)).fetchone()
    if not gasto: flash('Gasto no encontrado.', 'error'); return redirect(url_for('pagos'))
    saldo_actual = gasto['saldo_pendiente'] if gasto['saldo_pendiente'] is not None else gasto['monto']
    if not monto or monto <= 0 or monto > saldo_actual:
        flash('Monto de abono inválido.', 'error'); return redirect(url_for('pagos'))
    nuevo_saldo = round(saldo_actual - monto, 2)
    nuevo_estado = 'Pagado' if nuevo_saldo <= 0 else 'Parcial'
    try:
        db.execute('INSERT INTO abonos_gastos (id_gasto, monto, metodo_pago, referencia_pago, saldo_actual_snapshot) VALUES (?, ?, ?, ?, ?)',
                   (id_gasto, monto, metodo, referencia, nuevo_saldo))
        if nuevo_estado == 'Pagado':
            db.execute('UPDATE gastos SET saldo_pendiente = ?, estado = ?, fecha_pago = ? WHERE id = ?',
                       (nuevo_saldo, nuevo_estado, date.today().strftime('%Y-%m-%d'), id_gasto))
        else:
            db.execute('UPDATE gastos SET saldo_pendiente = ?, estado = ? WHERE id = ?', (nuevo_saldo, nuevo_estado, id_gasto))
        db.commit(); flash('Abono registrado.', 'success')
    except Exception as e:
        db.rollback(); flash(f'Error: {e}', 'error')
    return redirect(url_for('pagos'))

@app.route('/api/pagos/<int:id_gasto>/abonos')
@login_required
def api_gasto_abonos(id_gasto):
    db = get_db()
    abonos = db.execute(
        "SELECT id, datetime(fecha_abono, 'localtime') as fecha_abono, monto, metodo_pago, referencia_pago, saldo_actual_snapshot "
        "FROM abonos_gastos WHERE id_gasto = ? ORDER BY fecha_abono DESC", (id_gasto,)
    ).fetchall()
    return jsonify([dict(row) for row in abonos])

@app.route('/pagos/eliminar/<int:id_gasto>', methods=['POST'])
@login_required
@admin_required
def eliminar_gasto(id_gasto):
    db = get_db(); gasto = db.execute("SELECT comprobante_ruta FROM gastos WHERE id = ?", (id_gasto,)).fetchone()
    if gasto and gasto['comprobante_ruta']:
        try: os.remove(os.path.join(app.config['EXPENSE_UPLOAD_FOLDER'], gasto['comprobante_ruta']))
        except: pass
    db.execute("DELETE FROM abonos_gastos WHERE id_gasto = ?", (id_gasto,))
    db.execute("DELETE FROM gastos WHERE id = ?", (id_gasto,)); db.commit(); flash('Gasto eliminado.', 'success')
    return redirect(url_for('pagos'))

# =============================================================================
# --- 8. PAPELERA Y ELIMINACIÓN PERMANENTE ---
# =============================================================================

@app.route('/papelera')
@login_required
@admin_required
def eliminados_recientemente():
    db = get_db()
    p_el = db.execute('SELECT * FROM pacientes WHERE is_deleted = 1 ORDER BY deleted_at DESC').fetchall()
    i_el = db.execute('SELECT * FROM inventario WHERE is_deleted = 1 ORDER BY deleted_at DESC').fetchall()
    return render_template('eliminados.html', active_page='eliminados_recientemente', pacientes=p_el, items_inventario=i_el)

@app.route('/pacientes/eliminar/<int:id>', methods=['POST'])
@login_required
@admin_required
def eliminar_paciente(id):
    db = get_db(); db.execute("UPDATE pacientes SET is_deleted = 1, deleted_at = CURRENT_TIMESTAMP WHERE id = ?", (id,)); db.commit()
    flash('Paciente movido a la papelera.', 'success'); return redirect(url_for('pacientes'))

@app.route('/pacientes/restaurar/<int:id>', methods=['POST'])
@login_required
@admin_required
def restaurar_paciente(id):
    db = get_db(); db.execute("UPDATE pacientes SET is_deleted = 0, deleted_at = NULL WHERE id = ?", (id,)); db.commit()
    flash('Paciente restaurado.', 'success'); return redirect(url_for('eliminados_recientemente'))

@app.route('/pacientes/eliminar-permanente/<int:id>', methods=['POST'])
@login_required
@admin_required
def eliminar_paciente_permanente(id):
    db = get_db(); db.execute("DELETE FROM pacientes WHERE id = ?", (id,)); db.commit()
    flash('Paciente eliminado permanentemente.', 'success'); return redirect(url_for('eliminados_recientemente'))

@app.route('/inventario/item/eliminar/<int:id>', methods=['POST'])
@login_required
@admin_required
def inventario_eliminar(id):
    db = get_db(); db.execute('UPDATE inventario SET is_deleted = 1, deleted_at = CURRENT_TIMESTAMP WHERE id = ?', (id,)); db.commit()
    flash('Ítem movido a papelera.', 'success'); return redirect(url_for('inventario'))

@app.route('/inventario/item/restaurar/<int:id>', methods=['POST'])
@login_required
@admin_required
def inventario_restaurar(id):
    db = get_db(); db.execute('UPDATE inventario SET is_deleted = 0, deleted_at = NULL WHERE id = ?', (id,)); db.commit()
    flash('Ítem restaurado.', 'success'); return redirect(url_for('eliminados_recientemente'))

@app.route('/inventario/item/eliminar-permanente/<int:id>', methods=['POST'])
@login_required
@admin_required
def inventario_eliminar_permanente(id):
    db = get_db(); it = db.execute("SELECT ruta_foto FROM inventario WHERE id = ?", (id,)).fetchone()
    if it and it['ruta_foto']:
        try: os.remove(os.path.join(app.config['PRODUCT_UPLOAD_FOLDER'], it['ruta_foto']))
        except: pass
    db.execute('DELETE FROM inventario WHERE id = ?', (id,)); db.commit()
    flash('Ítem borrado permanentemente.', 'success'); return redirect(url_for('eliminados_recientemente'))

@app.route('/documento/eliminar/<int:id_documento>', methods=['POST'])
@login_required
@admin_required
def eliminar_documento(id_documento):
    db = get_db(); doc = db.execute('SELECT id_expediente FROM documentos WHERE id_documento = ?', (id_documento,)).fetchone()
    if doc: db.execute('UPDATE documentos SET is_deleted = 1, deleted_at = CURRENT_TIMESTAMP WHERE id_documento = ?', (id_documento,)); db.commit()
    return redirect(url_for('expediente_editar', id_expediente=doc['id_expediente']) + '#documentos')

@app.route('/expediente/<int:id_expediente>/documentos_eliminados')
@login_required
@admin_required
def documentos_eliminados(id_expediente):
    db = get_db(); expediente = db.execute('SELECT * FROM expedientes WHERE id_expediente = ?', (id_expediente,)).fetchone()
    paciente = db.execute('SELECT * FROM pacientes WHERE id = ?', (expediente['id_paciente'],)).fetchone()
    docs = db.execute('SELECT * FROM documentos WHERE id_expediente = ? AND is_deleted = 1 ORDER BY deleted_at DESC', (id_expediente,)).fetchall()
    return render_template('documentos_eliminados.html', paciente=paciente, expediente=expediente, documentos=docs)

@app.route('/documento/restaurar/<int:id_documento>', methods=['POST'])
@login_required
@admin_required
def restaurar_documento(id_documento):
    db = get_db(); doc = db.execute('SELECT id_expediente FROM documentos WHERE id_documento = ?', (id_documento,)).fetchone()
    if doc: db.execute('UPDATE documentos SET is_deleted = 0, deleted_at = NULL WHERE id_documento = ?', (id_documento,)); db.commit()
    return redirect(url_for('documentos_eliminados', id_expediente=doc['id_expediente']))

@app.route('/documento/eliminar-permanente/<int:id_documento>', methods=['POST'])
@login_required
@admin_required
def eliminar_documento_permanente(id_documento):
    db = get_db(); doc = db.execute('SELECT id_expediente, ruta_archivo FROM documentos WHERE id_documento = ?', (id_documento,)).fetchone()
    if doc:
        try: os.remove(os.path.join(app.config['UPLOAD_FOLDER'], doc['ruta_archivo']))
        except: pass
        db.execute('DELETE FROM documentos WHERE id_documento = ?', (id_documento,)); db.commit()
        return redirect(url_for('documentos_eliminados', id_expediente=doc['id_expediente']))
    return redirect(url_for('pacientes'))

if __name__ == '__main__':
    for folder in [UPLOAD_FOLDER, PRODUCT_UPLOAD_FOLDER, EXPENSE_UPLOAD_FOLDER]:
        if not os.path.exists(folder): os.makedirs(folder)
    app.run(debug=True)
